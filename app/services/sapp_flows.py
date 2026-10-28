"""Extract MTP flow workbooks and store directional hourly schedules."""

import hashlib
import math
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from app.db.database import get_db
from sapp_scraper import SappExtractionJob, parse_delivery_date_value, store_records_in_database

FLOW_DATASET_ID = "mtp_flows"
FLOW_COLLECTION = "sapp_flows"
BATCH_COLLECTION = "sapp_network_upload_batches"
FLOW_KEY = ("delivery_date", "hour", "product", "interface_id")


def _product(sheet_name: str, title: str) -> str:
    name = re.sub(r"[^A-Z0-9]+", "", f"{sheet_name} {title}".upper())
    for token, product in (("BILATERAL", "BILATERAL"), ("FPMM", "FPM_M"),
                           ("FPMW", "FPM_W"), ("DAM", "DAM"), ("NET", "NET_TOTAL")):
        if token in name:
            return product
    raise ValueError(f"Unrecognized flow sheet: {sheet_name}")


def extract_flows(file_path: Path, job: SappExtractionJob) -> list[dict]:
    """Validate all sheets before returning any records; never infer missing products."""
    batch_id = hashlib.sha256(file_path.read_bytes()).hexdigest()
    records = []
    keys = set()
    dates = set()
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            if not any(value is not None for row in rows for value in row):
                continue
            delivery_date = None
            header_index = None
            title = " ".join(str(value) for row in rows[:11] for value in row
                             if value is not None)
            product = _product(sheet.title, title)
            for index, row in enumerate(rows):
                for column, value in enumerate(row):
                    if str(value).strip().lower().rstrip(":") == "delivery date":
                        delivery_date = next((parsed for item in row[column + 1:]
                                              if (parsed := parse_delivery_date_value(item))), None)
                    if str(value).strip().lower() == "hour":
                        header_index, hour_column = index, column
                if header_index is not None:
                    break
            if delivery_date is None or header_index is None:
                raise ValueError(f"Missing delivery date or Hour header in {sheet.title}")
            dates.add(delivery_date.isoformat())
            routes = []
            for column, value in enumerate(rows[header_index]):
                if column <= hour_column or value is None or not str(value).strip():
                    continue
                interface = str(value).strip()
                if "_TO_" not in interface:
                    raise ValueError(f"Invalid flow corridor {interface!r}")
                from_area, to_area = interface.split("_TO_", 1)
                if not from_area or not to_area:
                    raise ValueError(f"Invalid flow corridor {interface!r}")
                routes.append((column, interface, from_area, to_area))
            if not routes:
                raise ValueError(f"No flow corridors in {sheet.title}")
            hours = set()
            for row in rows[header_index + 1:]:
                label = str(row[hour_column] or "").strip()
                if not label or label.lower() in {"min", "max", "avg", "tot", "total"}:
                    continue
                match = re.fullmatch(r"(\d{1,2})\s*-\s*(\d{1,2})", label)
                # Some exports use 23-24; others wrap the final interval to 23-00.
                if not match:
                    raise ValueError(f"Invalid flow interval {label!r} in {sheet.title}")
                start, end = int(match[1]), int(match[2])
                if not 0 <= start <= 23 or end not in {start + 1, (start + 1) % 24}:
                    raise ValueError(f"Invalid flow interval {label!r} in {sheet.title}")
                hour = int(match[1]) + 1
                if not 1 <= hour <= 24 or hour in hours:
                    raise ValueError(f"Invalid or duplicate flow hour {label!r}")
                hours.add(hour)
                for column, interface, from_area, to_area in routes:
                    value = row[column] if column < len(row) else None
                    if value is None or isinstance(value, bool):
                        raise ValueError(f"Missing flow for {sheet.title}/{label}/{interface}")
                    try:
                        flow = float(str(value).replace(",", ""))
                    except (TypeError, ValueError) as exc:
                        raise ValueError(f"Invalid flow for {sheet.title}/{label}/{interface}") from exc
                    if not math.isfinite(flow):
                        raise ValueError("Flow must be a finite number")
                    key = (delivery_date.isoformat(), hour, product, interface)
                    if key in keys:
                        raise ValueError(f"Duplicate flow key {key}")
                    keys.add(key)
                    records.append(dict(zip(FLOW_KEY, key), hour_label=label,
                                        from_area=from_area, to_area=to_area, flow_mw=flow,
                                        batch_id=batch_id, source_file=file_path.name,
                                        metadata={"data_source": job.data_source,
                                                  "sheet": sheet.title}))
            if hours != set(range(1, 25)):
                raise ValueError(f"Expected all 24 hours in {sheet.title}")
        if not records or len(dates) != 1:
            raise ValueError("Expected a nonempty flow workbook for one delivery date")
    finally:
        workbook.close()
    return records


FLOWS_JOB = SappExtractionJob(
    name="mtp_flows", subject_template="MTP - DAM Flows for {delivery_date}",
    data_source="SAPP_MTP_FLOWS", collection_name=FLOW_COLLECTION,
    unique_key_fields=FLOW_KEY, extractor=extract_flows,
)


def store_flows(records: list[dict]) -> dict:
    db = get_db()
    db[FLOW_COLLECTION].create_index([(key, 1) for key in FLOW_KEY], unique=True,
                                     name="idx_sapp_flows_natural_key")
    db[BATCH_COLLECTION].create_index("batch_id", unique=True)
    result = store_records_in_database(records, FLOWS_JOB)
    now = datetime.now(timezone.utc)
    for batch_id in sorted({record["batch_id"] for record in records}):
        batch_records = [record for record in records if record["batch_id"] == batch_id]
        db[BATCH_COLLECTION].update_one({"batch_id": batch_id}, {
            "$set": {"source_type": FLOW_DATASET_ID, "delivery_date": batch_records[0]["delivery_date"],
                     "source_file": batch_records[0]["source_file"], "status": "success",
                     "total": len(batch_records),
                     "products": sorted({record["product"] for record in batch_records}),
                     "updated_at": now},
            "$setOnInsert": {"created_at": now},
        }, upsert=True)
    return {**result, "total": len(records),
            "products": sorted({record["product"] for record in records}),
            "batch_ids": sorted({record["batch_id"] for record in records})}
