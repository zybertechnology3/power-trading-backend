"""MTP DAM available transmission capacity workbook ingestion."""

import math
import re
from pathlib import Path

from openpyxl import load_workbook

from app.db.database import get_db
from sapp_scraper import SappExtractionJob, parse_delivery_date_value, store_records_in_database

ATC_DATASET_ID = "mtp_atc"
ATC_COLLECTION = "sapp_bm_atc_records"
ATC_KEY = ("delivery_date", "market", "row_index")
SUMMARY_ROWS = {"min": ("Min", 25), "max": ("Max", 26),
                "avg": ("Avg", 27), "tot": ("Tot", 28), "total": ("Tot", 28)}


def _row_identity(value) -> tuple[int | None, str, int, bool]:
    label = str(value).strip()
    summary = SUMMARY_ROWS.get(label.lower())
    if summary:
        return None, summary[0], summary[1], True
    match = re.fullmatch(r"(\d{1,2})\s*-\s*(\d{1,2})", label)
    if match:
        start, end = int(match[1]), int(match[2])
        if 0 <= start <= 23 and end in {start + 1, (start + 1) % 24}:
            return start + 1, label, start + 1, False
    elif re.fullmatch(r"\d{1,2}(?:\.0)?", label):
        hour = int(float(label))
        if 1 <= hour <= 24:
            return hour, label, hour, False
    raise ValueError(f"Invalid ATC hour or summary label: {label!r}")


def _capacity(value) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("ATC capacity must be numeric")
    try:
        number = float(str(value).replace(",", ""))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid ATC capacity: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError("ATC capacity must be finite")
    return number


def extract_atc(file_path: Path, job: SappExtractionJob) -> list[dict]:
    """Preserve dynamic corridor headers, nulls, and any supplied summary rows."""
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    records = []
    keys = set()
    dates = set()
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            if not any(value is not None for row in rows for value in row):
                continue
            delivery_date = None
            header_index = None
            for index, row in enumerate(rows):
                for column, value in enumerate(row):
                    if str(value).strip().lower().rstrip(":") == "delivery date":
                        delivery_date = next((parsed for item in row[column + 1:]
                                              if (parsed := parse_delivery_date_value(item))), None)
                    if str(value).strip().lower() == "hour":
                        header_index, hour_column = index, column
                if header_index is not None:
                    break
            if delivery_date is None or header_index is None:
                raise ValueError(f"Missing delivery date or Hour header in ATC sheet {sheet.title}")
            dates.add(delivery_date.isoformat())
            routes = [(column, str(value).strip())
                      for column, value in enumerate(rows[header_index])
                      if column > hour_column and value is not None and str(value).strip()]
            columns = [name for _, name in routes]
            if not columns or len(columns) != len(set(columns)):
                raise ValueError("ATC corridor headers must be nonempty and unique")
            hours = set()
            for row in rows[header_index + 1:]:
                value = row[hour_column]
                if value is None or not str(value).strip():
                    if any(row[column] is not None for column, _ in routes):
                        raise ValueError("ATC values found without an hour label")
                    continue
                hour, label, row_index, is_summary = _row_identity(value)
                key = (delivery_date.isoformat(), "BM_ATC", row_index)
                if key in keys:
                    raise ValueError(f"Duplicate ATC row: {key}")
                keys.add(key)
                if hour is not None:
                    hours.add(hour)
                records.append(dict(
                    zip(ATC_KEY, key), source_market="DAM", hour=hour,
                    hour_label=label, product=label, is_summary=is_summary,
                    columns=columns,
                    values={name: _capacity(row[column]) for column, name in routes},
                    source_file=file_path.name,
                    metadata={"data_source": job.data_source, "sheet": sheet.title},
                ))
            if hours != set(range(1, 25)):
                raise ValueError(f"Expected all 24 ATC hours in {sheet.title}")
        if not records or len(dates) != 1:
            raise ValueError("Expected a nonempty ATC workbook for one delivery date")
    finally:
        workbook.close()
    return records


ATC_JOB = SappExtractionJob(
    name="mtp_atc",
    subject_template="MTP - DAM - Available Transmission Capacities for {delivery_date}",
    data_source="SAPP_MTP_DAM_AVAILABLE_TRANSMISSION_CAPACITIES",
    collection_name=ATC_COLLECTION, unique_key_fields=ATC_KEY, extractor=extract_atc,
)


def store_atc(records: list[dict]) -> dict:
    get_db()[ATC_COLLECTION].create_index(
        [(key, 1) for key in ATC_KEY], unique=True, name="idx_sapp_atc_natural_key",
    )
    result = store_records_in_database(records, ATC_JOB)
    return {**result, "total": len(records), "market": "BM_ATC", "source_market": "DAM",
            "columns": records[0]["columns"],
            "hourly_rows": sum(not record["is_summary"] for record in records),
            "summary_rows": sum(record["is_summary"] for record in records)}
