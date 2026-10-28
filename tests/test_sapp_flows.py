"""Offline flow ingestion, signed sync, and GET API regression tests."""

import hashlib
import hmac
import json
import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.sapp_flows import router
from app.services.sapp_flows import FLOWS_JOB, extract_flows
from sapp_email_document_importer import CreditNoteEmailImporter, _job_for_attachment


class Cursor(list):
    def sort(self, fields):
        return Cursor(sorted(self, key=lambda row: tuple(row[key] for key, _ in fields)))


class Collection:
    def __init__(self):
        self.rows = []

    def create_index(self, *args, **kwargs):
        pass

    def update_one(self, query, update, upsert=False):
        row = next((row for row in self.rows if all(row.get(k) == v for k, v in query.items())), None)
        inserted = row is None
        if inserted:
            row = dict(query)
            self.rows.append(row)
            row.update(update.get('$setOnInsert', {}))
        row.update(update.get('$set', {}))
        return SimpleNamespace(upserted_id='new' if inserted else None, matched_count=int(not inserted))

    def bulk_write(self, operations, ordered=False):
        results = [self.update_one(op._filter, op._doc, upsert=True) for op in operations]
        return SimpleNamespace(upserted_count=sum(r.upserted_id is not None for r in results),
                               matched_count=sum(r.matched_count for r in results))

    def find(self, query, projection=None):
        def matches(row):
            for key, value in query.items():
                if isinstance(value, dict):
                    if not value['$gte'] <= row[key] <= value['$lte']:
                        return False
                elif row.get(key) != value:
                    return False
            return True
        return Cursor(dict(row) for row in self.rows if matches(row))


class Database(dict):
    def __missing__(self, key):
        self[key] = Collection()
        return self[key]


class FlowsTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'MTP-Flows_2026-10-28_1.xlsx'
        self.make_workbook()
        self.db = Database()
        for target in ('app.services.sapp_flows.get_db', 'app.db.database.get_db',
                       'app.services.sapp_sync_notifier.get_db', 'app.api.sapp_flows.get_db'):
            patcher = patch(target, return_value=self.db)
            patcher.start()
            self.addCleanup(patcher.stop)
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)

    def make_workbook(self, missing_hour=False, blank=False, bad_date=False):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in ('NET Market Flows', 'FPM-M Flows', 'FPM-W Flows', 'DAM Flows'):
            sheet = workbook.create_sheet(name)
            sheet.append(['Delivery Date:', '2026-10-29' if bad_date else '2026-10-28'])
            sheet.append(['Hour', 'BOT_TO_RSAN', 'RSAN_TO_BOT'])
            for hour in range(24 if not missing_hour else 23):
                sheet.append([f'{hour:02d}-{(hour+1)%24:02d}', None if blank else -93, 93])
            sheet.append(['Min', -93, 93])
        workbook.save(self.path)
        workbook.close()

    def message(self):
        message = EmailMessage()
        message['Subject'] = 'FW: MTP - DAM Flows for 2026/10/28 - From: MTP Admin'
        message.set_content('Forwarded flows')
        message.add_attachment(self.path.read_bytes(), maintype='application',
                               subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               filename=self.path.name)
        return message

    def test_parser_products_directions_and_midnight(self):
        rows = extract_flows(self.path, FLOWS_JOB)
        self.assertEqual(len(rows), 192)
        self.assertEqual({r['product'] for r in rows}, {'NET_TOTAL', 'DAM', 'FPM_W', 'FPM_M'})
        self.assertEqual({r['hour'] for r in rows}, set(range(1, 25)))
        self.assertEqual(rows[0]['flow_mw'], -93)
        self.assertEqual(rows[0]['from_area'], 'BOT')
        self.assertEqual(rows[46]['hour_label'], '23-00')
        self.assertEqual(_job_for_attachment(self.path.name), FLOWS_JOB)
        self.assertIsNone(_job_for_attachment('MTP-DAM-Available-Transmission-Capacities_2026-10-28_1.xlsx'))

    def test_incomplete_or_blank_workbook_rejected(self):
        for kwargs in ({'missing_hour': True}, {'blank': True}):
            self.make_workbook(**kwargs)
            with self.assertRaises(ValueError):
                extract_flows(self.path, FLOWS_JOB)

    @patch.dict('os.environ', {'SAPP_SYNC_NOTIFY_URL': 'https://frontend.example/sync',
                              'SAPP_SYNC_NOTIFY_SECRET': 'test-secret'})
    @patch('app.services.sapp_sync_notifier.httpx.post')
    def test_forwarded_email_storage_webhook_api_and_reimport(self, post):
        post.return_value.status_code = 200
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        result = importer._process_message('123', self.message())
        self.assertEqual(result['jobs'][0]['imported'], 192)
        self.assertEqual(len(self.db['sapp_network_upload_batches'].rows), 1)
        request = post.call_args.kwargs
        payload = json.loads(request['content'])
        self.assertEqual(payload['dataset_id'], 'mtp_flows')
        self.assertEqual(payload['event'], 'sapp.dataset.ready')
        self.assertEqual(payload['date_range']['start_date'], '2026-10-28')
        expected = 'sha256=' + hmac.new(b'test-secret', request['content'], hashlib.sha256).hexdigest()
        self.assertEqual(request['headers']['X-SAPP-Signature'], expected)
        importer._process_message('124', self.message())
        self.assertEqual(len(self.db['sapp_flows'].rows), 192)
        self.assertEqual(len(self.db['sapp_network_upload_batches'].rows), 1)
        self.assertEqual(len(self.db['sapp_auto_scraper_runs'].rows), 2)
        response = self.client.get('/sapp/flows/2026-10-28?product=DAM&from_area=BOT')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['total'], 24)
        self.assertEqual(response.json()['records'][0]['flow_mw'], -93)
        response = self.client.get('/sapp/flows?start_date=2026-10-28&end_date=2026-10-29')
        self.assertEqual(response.json()['total'], 192)
        self.assertEqual(self.client.get('/sapp/flows/2026-10-29').json()['total'], 0)

    @patch('sapp_email_document_importer.record_external_run')
    def test_date_mismatch_never_stored_or_notified(self, notify):
        self.make_workbook(bad_date=True)
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        with self.assertRaises(ValueError):
            importer._process_message('123', self.message())
        self.assertEqual(len(self.db), 0)
        notify.assert_not_called()

    @patch('sapp_email_document_importer.record_external_run', return_value={'status': 'sent'})
    def test_revised_workbook_updates_existing_key(self, notify):
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        importer._process_message('123', self.message())
        workbook = load_workbook(self.path)
        workbook['DAM Flows']['B3'] = -100
        workbook.save(self.path)
        workbook.close()
        importer._process_message('124', self.message())
        self.assertEqual(len(self.db['sapp_flows'].rows), 192)
        self.assertEqual(len(self.db['sapp_network_upload_batches'].rows), 2)
        response = self.client.get('/sapp/flows/2026-10-28?product=DAM&interface_id=BOT_TO_RSAN')
        self.assertEqual(response.json()['records'][0]['flow_mw'], -100)

    @patch('app.services.sapp_sync_notifier.httpx.post', side_effect=RuntimeError('offline'))
    @patch.dict('os.environ', {'SAPP_SYNC_NOTIFY_URL': 'https://frontend.example/sync',
                              'SAPP_SYNC_NOTIFY_SECRET': 'test-secret'})
    def test_webhook_failure_keeps_data_and_finished_run(self, post):
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        result = importer._process_message('123', self.message())
        self.assertEqual(result['jobs'][0]['notification']['status'], 'failed')
        self.assertEqual(len(self.db['sapp_flows'].rows), 192)
        self.assertEqual(self.db['sapp_auto_scraper_runs'].rows[0]['status'], 'success')

    def test_api_rejects_invalid_queries(self):
        for path, status in (('/sapp/flows', 400),
                             ('/sapp/flows?delivery_date=2026-10-28&start_date=2026-10-28', 400),
                             ('/sapp/flows?start_date=2026-10-29&end_date=2026-10-28', 400),
                             ('/sapp/flows/2026-10-28?product=invalid', 422)):
            self.assertEqual(self.client.get(path).status_code, status)


if __name__ == '__main__':
    unittest.main()
