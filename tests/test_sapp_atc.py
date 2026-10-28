"""Offline ATC extraction, email sync, and API regression tests."""

import hashlib
import hmac
import json
import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.sapp_atc import router
from app.services.sapp_atc import ATC_JOB, extract_atc
from sapp_email_document_importer import (
    CreditNoteEmailImporter, DEFAULT_SUBJECT_SEARCH_TERMS, _job_for_attachment,
)
from tests.test_sapp_flows import Database


class AtcTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'MTP-DAM-Available-Transmission-Capacities_2026-10-28_1.xlsx'
        self.make_workbook()
        self.db = Database()
        for target in ('app.services.sapp_atc.get_db', 'app.db.database.get_db',
                       'app.services.sapp_sync_notifier.get_db', 'app.api.sapp_atc.get_db'):
            patcher = patch(target, return_value=self.db)
            patcher.start()
            self.addCleanup(patcher.stop)
        app = FastAPI()
        app.include_router(router)
        self.client = TestClient(app)

    def make_workbook(self, summaries=True):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'ATC'
        sheet.append([None, 'Delivery Date:', '2026-10-28'])
        sheet.append([None, 'Hour', 'BOT_TO_RSAN', 'RSAN_TO_BOT'])
        for hour in range(24):
            sheet.append([None, f'{hour:02d}-{(hour+1)%24:02d}', 569.5, 0])
        if summaries:
            for label in ('Min', 'Max', 'Avg', 'Tot'):
                sheet.append([None, label, 569.5, None])
        workbook.save(self.path)
        workbook.close()

    def message(self):
        message = EmailMessage()
        message['Subject'] = ('FW: MTP - DAM - Available Transmission Capacities for '
                              '2026/10/28 - From: MTP Admin')
        message.set_content('Forwarded capacity workbook')
        message.add_attachment(self.path.read_bytes(), maintype='application',
                               subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                               filename=self.path.name)
        return message

    def test_hourly_rows_and_optional_summaries(self):
        rows = extract_atc(self.path, ATC_JOB)
        self.assertEqual(len(rows), 28)
        self.assertEqual(rows[0]['hour'], 1)
        self.assertEqual(rows[23]['hour'], 24)
        self.assertEqual(rows[23]['hour_label'], '23-00')
        self.assertEqual(rows[0]['columns'], ['BOT_TO_RSAN', 'RSAN_TO_BOT'])
        self.assertEqual(rows[0]['values']['RSAN_TO_BOT'], 0)
        for index, label in enumerate(('Min', 'Max', 'Avg', 'Tot'), 24):
            self.assertEqual(rows[index]['hour_label'], label)
            self.assertEqual(rows[index]['row_index'], index + 1)
            self.assertIsNone(rows[index]['hour'])
            self.assertTrue(rows[index]['is_summary'])
            self.assertIsNone(rows[index]['values']['RSAN_TO_BOT'])
        self.make_workbook(summaries=False)
        self.assertEqual(len(extract_atc(self.path, ATC_JOB)), 24)

    def test_subject_and_filename_classification(self):
        self.assertIn('MTP - DAM - Available Transmission Capacities', DEFAULT_SUBJECT_SEARCH_TERMS)
        self.assertEqual(_job_for_attachment(self.path.name), ATC_JOB)
        self.assertEqual(_job_for_attachment(self.path.stem + ' (1).xlsx'), ATC_JOB)
        self.assertIsNone(_job_for_attachment('MTP-FPM-W-Available-Transmission-Capacities_2026-10-28.xlsx'))

    def test_invalid_workbook_rejected(self):
        for mutation in ('missing_hour', 'duplicate_hour', 'bad_value', 'duplicate_header', 'bad_interval'):
            self.make_workbook()
            workbook = load_workbook(self.path)
            sheet = workbook.active
            if mutation == 'missing_hour':
                sheet.delete_rows(3)
            elif mutation == 'duplicate_hour':
                sheet['B4'] = '00-01'
            elif mutation == 'bad_value':
                sheet['C3'] = 'not-a-number'
            elif mutation == 'duplicate_header':
                sheet['D2'] = 'BOT_TO_RSAN'
            else:
                sheet['B3'] = '00-03'
            workbook.save(self.path)
            workbook.close()
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                extract_atc(self.path, ATC_JOB)

    @patch.dict('os.environ', {'SAPP_SYNC_NOTIFY_URL': 'https://frontend.example/sync',
                              'SAPP_SYNC_NOTIFY_SECRET': 'test-secret'})
    @patch('app.services.sapp_sync_notifier.httpx.post')
    def test_forwarded_email_signed_webhook_and_api(self, post):
        post.return_value.status_code = 200
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        result = importer._process_message('123', self.message())
        self.assertEqual(result['jobs'][0]['imported'], 28)
        request = post.call_args.kwargs
        payload = json.loads(request['content'])
        self.assertEqual(payload['dataset_id'], 'mtp_atc')
        self.assertEqual(payload['job'], 'mtp_atc')
        self.assertEqual(payload['event_id'], 'email:123:mtp_atc')
        self.assertEqual(payload['event'], 'sapp.dataset.ready')
        self.assertEqual(payload['date_range']['start_date'], '2026-10-28')
        expected = 'sha256=' + hmac.new(b'test-secret', request['content'], hashlib.sha256).hexdigest()
        self.assertEqual(request['headers']['X-SAPP-Signature'], expected)
        importer._process_message('124', self.message())
        self.assertEqual(len(self.db['sapp_bm_atc_records'].rows), 28)
        response = self.client.get('/sapp/atc/2026-10-28')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['dataset_id'], 'mtp_atc')
        self.assertEqual(response.json()['total'], 28)
        self.assertEqual(response.json()['records'][0]['source_market'], 'DAM')
        self.assertEqual(response.json()['records'][-1]['hour_label'], 'Tot')
        self.assertEqual(self.client.get('/sapp/atc/2026-10-28?include_summaries=false').json()['total'], 24)
        self.assertEqual(self.client.get('/sapp/atc?start_date=2026-10-28&end_date=2026-10-29').json()['total'], 28)
        self.assertEqual(self.client.get('/sapp/atc?delivery_date=2026-10-29').json()['total'], 0)

    @patch('sapp_email_document_importer.record_external_run', return_value={'status': 'sent'})
    def test_revised_workbook_updates_existing_values(self, notify):
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        importer._process_message('123', self.message())
        workbook = load_workbook(self.path)
        workbook.active['C3'] = 600
        workbook.save(self.path)
        workbook.close()
        importer._process_message('124', self.message())
        self.assertEqual(len(self.db['sapp_bm_atc_records'].rows), 28)
        rows = self.client.get('/sapp/atc/2026-10-28').json()['records']
        self.assertEqual(rows[0]['values']['BOT_TO_RSAN'], 600)

    @patch('sapp_email_document_importer.record_external_run')
    def test_date_mismatch_never_stored_or_notified(self, notify):
        workbook = load_workbook(self.path)
        workbook.active['C1'] = '2026-10-29'
        workbook.save(self.path)
        workbook.close()
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        with self.assertRaises(ValueError):
            importer._process_message('123', self.message())
        self.assertEqual(len(self.db), 0)
        notify.assert_not_called()

    @patch('app.services.sapp_sync_notifier.httpx.post', side_effect=RuntimeError('offline'))
    @patch.dict('os.environ', {'SAPP_SYNC_NOTIFY_URL': 'https://frontend.example/sync',
                              'SAPP_SYNC_NOTIFY_SECRET': 'test-secret'})
    def test_webhook_failure_retains_data_and_completion(self, post):
        importer = CreditNoteEmailImporter.__new__(CreditNoteEmailImporter)
        result = importer._process_message('123', self.message())
        self.assertEqual(result['jobs'][0]['notification']['status'], 'failed')
        self.assertEqual(len(self.db['sapp_bm_atc_records'].rows), 28)
        self.assertEqual(self.db['sapp_auto_scraper_runs'].rows[0]['dataset_id'], 'mtp_atc')

    def test_invalid_api_queries(self):
        for path, status in (('/sapp/atc', 400),
                             ('/sapp/atc?delivery_date=2026-10-28&end_date=2026-10-28', 400),
                             ('/sapp/atc?start_date=2026-10-29&end_date=2026-10-28', 400),
                             ('/sapp/atc/not-a-date', 422)):
            self.assertEqual(self.client.get(path).status_code, status)


if __name__ == '__main__':
    unittest.main()
